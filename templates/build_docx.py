"""Markdown -> formatted .docx (Harvard-friendly). Usage:
venv/bin/python -I templates/build_docx.py draft.md out.docx --code ABC1234 --title "Title" [--toc] [--no-title-page] [--base-dir DIR]
    [--cover --school "..." --module-title "..." --sid "..." --year 2025-26 --assessment-title "..."]
--cover prepends a generic UK-style Coursework Cover Sheet (adapt the COVER_* texts to your university). Only the
department/assessment fields are filled; the declaration tick-boxes, the AI-tools disclosure, well-being and
reflection questions are left blank for the student to complete personally.
Supports: # ## ### headings, paragraphs, **bold**, *italic*, - bullets, 1. lists,
![Figure 1: caption](path.png) (path relative to the .md or --base-dir), a 'Table N: caption' line directly above a | table | (caption above table; placeholders [USER DATA]/[UNVERIFIED]/[TODO] are highlighted), and a '## References' section (hanging indent).
"""
import argparse, datetime, os, re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def field(par, instr):
    for t, txt in (("begin", None), (None, instr), ("end", None)):
        r = par.add_run()
        if t:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        r._r.append(e)

PLACEHOLDER = re.compile(r"(\[(?:USER DATA|UNVERIFIED|TODO)[^\]]*\])")

def plain(par, text, **fmt):
    """Add text; any [USER DATA: ...]/[UNVERIFIED]/[TODO] marker is kept verbatim but highlighted yellow."""
    for seg in PLACEHOLDER.split(text):
        if not seg: continue
        r = par.add_run(seg)
        for k, v in fmt.items(): setattr(r, k, v)
        if PLACEHOLDER.fullmatch(seg): r.font.highlight_color = WD_COLOR_INDEX.YELLOW

def inline(par, text):
    for part in re.split(r"(\*\*[^*]+\*\*|\*[^*\s][^*]*\*)", text):
        if part.startswith("**") and part.endswith("**") and len(part) > 4: plain(par, part[2:-2], bold=True)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: plain(par, part[1:-1], italic=True)
        elif part: plain(par, part)

COVER_INTRO = ("Students should complete the input fields contained in this form and attach it in front of your formal "
               "assessment submission. All fields within this form are required. Please ensure that check boxes and "
               "radio buttons are appropriately selected. The last three questions are just for you to personally consider.")
COVER_DECLARATION = ("By attaching this front cover sheet to my assessment I confirm and declare that I am the sole author "
                     "of this work, except where otherwise acknowledged by appropriate referencing and citation, and that "
                     "I have taken all reasonable skill and care to ensure that no other person has been able, or allowed, "
                     "to copy this work in either paper or electronic form, and that prior to submission I have read, "
                     "understood and followed the University regulations as outlined in the Academic Integrity Policy and "
                     "Procedure for Academic Misconduct")
COVER_CHECKS = ("I have applied the learning outcomes for this module",
                "I have checked for Academic Integrity via Turn-it-in",
                "I have used references in accordance with instructions in the Assessment Brief",
                "I have proofread my work for spelling, grammar and punctuation.",
                "I have checked that the word count/size of this submissions accords with the guidance provided in the Assessment Brief.",
                "I have followed the guidance in the Assessment Brief and have not made inappropriate use of AI to boost my grade unfairly.")
COVER_WELLBEING = ("We wish to support any student who is experiencing mitigating circumstances which prevents students from "
                   "performing to the best of their ability when completing or submitting assignments. If you are experiencing "
                   "such circumstances, then you may apply for Mitigating Circumstances. Wherever possible this must be done "
                   "prior to handing in your assignment.")

def cover_sheet(d, a):
    def para(text="", bold=False, size=11, after=4):
        p = d.add_paragraph(); p.paragraph_format.line_spacing = 1.1; p.paragraph_format.space_after = Pt(after)
        if text:
            r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        return p
    def field_line(label, value):
        p = para(); r = p.add_run(label + " "); r.bold = True
        p.add_run(value if value else "____________________")
    para("Coursework Cover Sheet", True, 18, 8)
    para(COVER_INTRO)
    para("Department and assessment information:", True, 12)
    field_line("School Name:", a.school)
    field_line("Assessment title:", a.assessment_title)
    field_line("Course Title:", a.course)
    field_line("Module Title:", a.module_title)
    field_line("Module Code:", a.code if not a.code.startswith("[") else "")
    field_line("Year of Study:", a.year)
    field_line("Student Number:", a.sid if not a.sid.startswith("[") else "")
    para("Academic Misconduct / Plagiarism Declaration", True, 12)
    para(COVER_DECLARATION)
    para("Have you checked the following? This will help your assessment achievement.", True, 12)
    for c in COVER_CHECKS: para("☐  " + c)
    para("List any AI tools used (including assistive tools such as Grammarly)", True)
    para("Please provide details of how you have used AI tools, including examples such as screenshots of prompts and the output which have been generated.")
    para("____________________________________________")
    para("Well-being", True, 12)
    para(COVER_WELLBEING)
    para("Do you need to apply for mitigating circumstances for this assignment? Please select:  ☐ Yes   ☐ No")
    para("Questions you may wish to consider:", True)
    para("1. Have I allowed sufficient time to prepare this assessment?  ☐ Yes  ☐ No")
    para("2. Have I reflected on previous feedback and made improvements in accordance with advice?  ☐ Yes  ☐ No")
    para("3. What grade am I expecting: ______")
    d.add_page_break()

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("md"); ap.add_argument("out")
    ap.add_argument("--code", default="[Module code]"); ap.add_argument("--title", default="Title")
    ap.add_argument("--name", default="[Student name]"); ap.add_argument("--sid", default="[Student ID]")
    ap.add_argument("--toc", action="store_true"); ap.add_argument("--no-title-page", action="store_true")
    ap.add_argument("--cover", action="store_true", help="prepend the Coursework Cover Sheet")
    ap.add_argument("--school", default=""); ap.add_argument("--course", default="")
    ap.add_argument("--module-title", default=""); ap.add_argument("--year", default="2025-26")
    ap.add_argument("--assessment-title", default=""); ap.add_argument("--date", default=datetime.date.today().strftime("%d %B %Y"))
    ap.add_argument("--base-dir", default=None, help="folder that relative image paths resolve against (default: folder of the .md)")
    a = ap.parse_args()
    d = Document()
    s = d.sections[0]
    for m in ("left_margin", "right_margin", "top_margin", "bottom_margin"): setattr(s, m, Cm(2.54))
    st = d.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(6)
    for n, sz in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        h = d.styles[n]; h.font.name = "Arial"; h.font.size = Pt(sz); h.font.bold = True
        h.font.color.rgb = None
        rf = h.element.rPr.rFonts
        for att in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            if rf.get(qn(att)) is not None: del rf.attrib[qn(att)]
        rf.set(qn("w:eastAsia"), "Arial"); rf.set(qn("w:cs"), "Arial")
        h.paragraph_format.keep_with_next = True
    cs = d.styles["Caption"]; cs.font.name = "Arial"; cs.font.size = Pt(10); cs.font.italic = True; cs.font.bold = False
    cs.font.color.rgb = None; cs.paragraph_format.line_spacing = 1.15; cs.paragraph_format.space_after = Pt(8)
    d.core_properties.title = a.title; d.core_properties.author = a.name if not a.name.startswith("[") else ""
    fp = s.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; field(fp, "PAGE")
    if a.cover: cover_sheet(d, a)
    if not a.no_title_page:
        for _ in range(6): d.add_paragraph()
        items = ((a.code, 14, False), (a.module_title, 14, False), (a.title, 22, True), (a.name, 12, False), (a.sid, 12, False), (a.date, 12, False))
        for txt, sz, b in items:
            if not txt or txt.startswith("["): continue  # never print unfilled placeholders
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(txt); r.font.size = Pt(sz); r.bold = b
        d.add_page_break()
    if a.toc:
        d.add_heading("Contents", 1); field(d.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u'); d.add_page_break()
    base = a.base_dir or os.path.dirname(os.path.abspath(a.md))
    lines = open(a.md, encoding="utf-8").read().splitlines()
    pending_caption = None; words = 0; n_fig = 0; n_tab = 0
    in_refs = False; i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln: i += 1; continue
        m = re.match(r"^(#{1,3})\s+(.*)", ln)
        if m:
            in_refs = m.group(2).strip().lower() in ("references", "reference list", "bibliography")
            d.add_heading(m.group(2), len(m.group(1))); i += 1; continue
        m = re.match(r"^!\[(.*?)\]\((.*?)\)", ln)
        if m:
            p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img = m.group(2) if os.path.isabs(m.group(2)) else os.path.join(base, m.group(2))
            if not os.path.exists(img): sys.exit("image not found: " + img)
            p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(img, width=Cm(14)); n_fig += 1
            c = d.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inline(c, m.group(1)); i += 1; continue
        m = re.match(r"^(?:\*\*)?(Table \d+[:.].*?)(?:\*\*)?$", ln)   # table caption line (goes above the table)
        nxt = next((j for j in range(i + 1, min(i + 3, len(lines))) if lines[j].strip()), None)  # allow one blank line
        if m and nxt is not None and lines[nxt].startswith("|"):
            c = d.add_paragraph(style="Caption"); c.alignment = WD_ALIGN_PARAGRAPH.CENTER
            c.paragraph_format.keep_with_next = True; c.paragraph_format.space_after = Pt(4)
            inline(c, m.group(1)); i += 1; continue
        if ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:|-]+\|?$", lines[i]): rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            t = d.add_table(rows=len(rows), cols=len(rows[0])); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
            n_tab += 1
            for ri, row in enumerate(rows):
                for ci, v in enumerate(row):
                    cell = t.cell(ri, ci); cell.text = ""; cp = cell.paragraphs[0]; inline(cp, v)
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.line_spacing = 1.0; cp.paragraph_format.space_after = Pt(2); cp.paragraph_format.space_before = Pt(2)
                    if ri == 0:
                        for r in cp.runs: r.bold = True
            if rows: 
                tr = t.rows[0]._tr.get_or_add_trPr(); th = OxmlElement("w:tblHeader"); tr.append(th)
            d.add_paragraph().paragraph_format.space_after = Pt(0); continue
        m = re.match(r"^\s*[-*]\s+(.*)", ln)
        if m: inline(d.add_paragraph(style="List Bullet"), m.group(1)); i += 1; continue
        m = re.match(r"^\s*\d+[.)]\s+(.*)", ln)
        if m: inline(d.add_paragraph(style="List Number"), m.group(1)); i += 1; continue
        p = d.add_paragraph(); inline(p, ln)
        if not in_refs: words += len(re.findall(r"\S+", ln))
        if in_refs:
            p.paragraph_format.left_indent = Cm(1.27); p.paragraph_format.first_line_indent = Cm(-1.27)
            p.paragraph_format.line_spacing = 1.15
        i += 1
    d.save(a.out); print("saved", a.out)
    print(f"body words (paragraphs only; excl. headings, tables, captions, references): {words}; figures: {n_fig}; tables: {n_tab}")

if __name__ == "__main__": main()
